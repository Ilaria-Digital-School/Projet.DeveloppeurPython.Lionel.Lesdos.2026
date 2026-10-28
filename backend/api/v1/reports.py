"""
API REST v1 — Rapports.

Routes:
    GET    /api/v1/rapports              Liste des rapports sauvegardés
    POST   /api/v1/rapports/generate     Générer un rapport
    GET    /api/v1/rapports/data         Données de rapport pour graphiques
    GET    /api/v1/rapports/<id>         Récupérer un rapport sauvegardé
    DELETE /api/v1/rapports/<id>         Supprimer un rapport sauvegardé
    GET    /api/v1/rapports/export       Exporter un rapport (pdf, excel, csv, json)
    GET    /api/v1/rapports/export/pdf   Exporter un rapport en PDF (alias)
"""

import csv
import json
import logging
from datetime import datetime, timedelta
from io import BytesIO, StringIO

from flask import make_response, request
from flask_login import current_user, login_required
from sqlalchemy import func

from core.extensions import db
from core.models.category import Category
from core.models.report import Report
from core.models.transactions import Transaction, TransactionType

from . import api_v1_bp
from .auth import auth_limiter
from .helpers import api_error, api_success

logger = logging.getLogger(__name__)


# =============================
# FONCTIONS UTILITAIRES PRIVEES
# =============================


def _resolve_date_range(date_range, params):
    """Résout une période en dates de début et fin."""
    today = datetime.now().date()

    if date_range == "current_month":
        start_date = today.replace(day=1)
        end_date = today
    elif date_range == "last_month":
        first_of_month = today.replace(day=1)
        end_date = first_of_month - timedelta(days=1)
        start_date = end_date.replace(day=1)
    elif date_range == "last_6_months":
        start_date = (today - timedelta(days=180)).replace(day=1)
        end_date = today
    elif date_range == "current_year":
        start_date = today.replace(month=1, day=1)
        end_date = today
    elif date_range == "last_year":
        start_date = today.replace(year=today.year - 1, month=1, day=1)
        end_date = today.replace(year=today.year - 1, month=12, day=31)
    elif date_range == "custom":
        start_str = params.get("start_date")
        end_str = params.get("end_date")
        start_date = (
            datetime.strptime(start_str, "%Y-%m-%d").date()
            if start_str
            else today.replace(day=1)
        )
        end_date = datetime.strptime(end_str, "%Y-%m-%d").date() if end_str else today
    else:
        start_date = today.replace(day=1)
        end_date = today

    return start_date, end_date


def _normalize_report_type(report_type):
    """Normalise le type de rapport (accepte underscore et tiret)."""
    return report_type.replace("_", "-") if report_type else "expenses"


def _generate_expenses_report(user_id, start_date, end_date):
    """Génère un rapport de dépenses par catégorie."""
    # Jointure externe : les dépenses sans catégorie sont comptées sous
    # "Non catégorisée" (sinon le total diverge du rapport revenus/dépenses)
    results = (
        db.session.query(
            func.coalesce(Category.name, "Non catégorisée").label("category"),
            func.sum(Transaction.montant).label("total"),
        )
        .select_from(Transaction)
        .outerjoin(Category, Transaction.category_id == Category.id)
        .filter(
            Transaction.user_id == user_id,
            func.date(Transaction.date_heure_transaction) >= start_date,
            func.date(Transaction.date_heure_transaction) <= end_date,
            Transaction.type_transaction.in_(TransactionType.expense_values()),
        )
        .group_by(Category.name)
        .order_by(func.sum(Transaction.montant).desc())
        .all()
    )

    categories = []
    amounts = []
    total = 0.0
    for row in results:
        amount = abs(float(row.total or 0))
        categories.append(row.category)
        amounts.append(amount)
        total += amount

    # Calculer les pourcentages
    data_items = []
    for cat, amt in zip(categories, amounts):
        pct = round((amt / total * 100), 1) if total > 0 else 0
        data_items.append({"category": cat, "amount": amt, "percentage": pct})

    return {
        "categories": categories,
        "amounts": amounts,
        "total": total,
        "items": data_items,
    }


def _generate_income_vs_expenses_report(user_id, start_date, end_date):
    """Génère un rapport revenus vs dépenses par mois."""
    expense_types = TransactionType.expense_values()
    income_types = TransactionType.income_values()

    # Dépenses par mois
    expenses = (
        db.session.query(
            func.date_trunc("month", Transaction.date_heure_transaction).label(
                "period"
            ),
            func.sum(Transaction.montant).label("total"),
        )
        .filter(
            Transaction.user_id == user_id,
            func.date(Transaction.date_heure_transaction) >= start_date,
            func.date(Transaction.date_heure_transaction) <= end_date,
            Transaction.type_transaction.in_(expense_types),
        )
        .group_by("period")
        .order_by("period")
        .all()
    )

    # Revenus par mois
    incomes = (
        db.session.query(
            func.date_trunc("month", Transaction.date_heure_transaction).label(
                "period"
            ),
            func.sum(Transaction.montant).label("total"),
        )
        .filter(
            Transaction.user_id == user_id,
            func.date(Transaction.date_heure_transaction) >= start_date,
            func.date(Transaction.date_heure_transaction) <= end_date,
            Transaction.type_transaction.in_(income_types),
        )
        .group_by("period")
        .order_by("period")
        .all()
    )

    expense_map = {str(r.period)[:7]: abs(float(r.total or 0)) for r in expenses}
    income_map = {str(r.period)[:7]: float(r.total or 0) for r in incomes}

    all_periods = sorted(set(list(expense_map.keys()) + list(income_map.keys())))

    periods = []
    expenses_data = []
    income_data = []
    for p in all_periods:
        periods.append(p)
        expenses_data.append(expense_map.get(p, 0))
        income_data.append(income_map.get(p, 0))

    total_income = sum(income_data)
    total_expenses = sum(expenses_data)

    return {
        "periods": periods,
        "expenses": expenses_data,
        "income": income_data,
        "total_income": total_income,
        "total_expenses": total_expenses,
        "savings": total_income - total_expenses,
    }


def _generate_cashflow_report(user_id, start_date, end_date):
    """Génère un rapport de flux de trésorerie."""
    income_data = _generate_income_vs_expenses_report(user_id, start_date, end_date)

    # Calculer le flux net cumulé
    cumulative = []
    running_total = 0
    for inc, exp in zip(income_data["income"], income_data["expenses"]):
        running_total += inc - exp
        cumulative.append(running_total)

    income_data["cumulative"] = cumulative
    return income_data


def _build_pdf_report(
    report_type, report_data, start_date, end_date, include_charts=False
):
    """Construit un document PDF pour un rapport."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
    )

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=72,
        leftMargin=72,
        topMargin=72,
        bottomMargin=18,
    )

    styles = getSampleStyleSheet()
    elements = []

    # Titre
    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Heading1"],
        fontSize=18,
        spaceAfter=20,
        alignment=1,
        textColor=colors.darkblue,
    )

    type_labels = {
        "expenses": "Rapport des Depenses par Categorie",
        "expenses_by_category": "Rapport des Depenses par Categorie",
        "expenses-by-category": "Rapport des Depenses par Categorie",
        "income-vs-expenses": "Rapport Revenus vs Depenses",
        "cashflow": "Rapport Flux de Tresorerie",
    }
    title_text = type_labels.get(report_type, "Rapport Financier")
    elements.append(Paragraph(title_text, title_style))
    elements.append(Spacer(1, 10))

    # Sous-titre avec période
    period_text = f"Periode : du {start_date.strftime('%d/%m/%Y')} au {end_date.strftime('%d/%m/%Y')}"
    elements.append(Paragraph(period_text, styles["Normal"]))
    elements.append(
        Paragraph(
            f"Genere le {datetime.now().strftime('%d/%m/%Y a %H:%M')}", styles["Normal"]
        )
    )
    elements.append(Spacer(1, 20))

    if report_type in ("expenses", "expenses_by_category", "expenses-by-category"):
        _add_expenses_pdf_content(elements, styles, report_data)
        if include_charts:
            _add_expenses_chart(elements, report_data)
    elif report_type in ("income-vs-expenses", "cashflow"):
        _add_income_vs_expenses_pdf_content(elements, styles, report_data)
        if include_charts:
            _add_income_vs_expenses_chart(elements, report_data)

    doc.build(elements)
    buffer.seek(0)
    return buffer


def _add_expenses_pdf_content(elements, styles, data):
    """Ajoute le contenu d'un rapport de dépenses au PDF."""
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

    total = data.get("total", 0)
    elements.append(
        Paragraph(f"Total des depenses : {total:.2f} EUR", styles["Heading2"])
    )
    elements.append(Spacer(1, 15))

    items = data.get("items", [])
    if items:
        table_data = [["Categorie", "Montant", "%"]]
        for item in items:
            table_data.append(
                [
                    item["category"],
                    f"{item['amount']:.2f} EUR",
                    f"{item['percentage']}%",
                ]
            )
        table_data.append(["TOTAL", f"{total:.2f} EUR", "100%"])

        table = Table(table_data, colWidths=[250, 120, 80])
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("ALIGN", (0, 1), (0, -1), "LEFT"),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 12),
                    ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 1), (-1, -1), 10),
                    ("GRID", (0, 0), (-1, -1), 1, colors.black),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.beige, colors.white]),
                    ("BACKGROUND", (0, -1), (-1, -1), colors.lightblue),
                    ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ]
            )
        )
        elements.append(table)
    else:
        elements.append(
            Paragraph("Aucune depense trouvee pour cette periode.", styles["Normal"])
        )


def _add_income_vs_expenses_pdf_content(elements, styles, data):
    """Ajoute le contenu d'un rapport revenus/dépenses au PDF."""
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

    total_income = data.get("total_income", 0)
    total_expenses = data.get("total_expenses", 0)
    savings = data.get("savings", 0)

    summary_style = styles["Heading2"]
    elements.append(
        Paragraph(f"Revenus totaux : {total_income:.2f} EUR", summary_style)
    )
    elements.append(
        Paragraph(f"Depenses totales : {total_expenses:.2f} EUR", summary_style)
    )
    elements.append(Paragraph(f"Economies : {savings:.2f} EUR", summary_style))
    elements.append(Spacer(1, 15))

    periods = data.get("periods", [])
    incomes = data.get("income", [])
    expenses_list = data.get("expenses", [])

    if periods:
        table_data = [["Periode", "Revenus", "Depenses", "Solde"]]
        for p, inc, exp in zip(periods, incomes, expenses_list):
            balance = inc - exp
            table_data.append(
                [p, f"{inc:.2f} EUR", f"{exp:.2f} EUR", f"{balance:.2f} EUR"]
            )
        table_data.append(
            [
                "TOTAL",
                f"{total_income:.2f} EUR",
                f"{total_expenses:.2f} EUR",
                f"{savings:.2f} EUR",
            ]
        )

        table = Table(table_data, colWidths=[100, 120, 120, 120])
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 12),
                    ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 1), (-1, -1), 10),
                    ("GRID", (0, 0), (-1, -1), 1, colors.black),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.beige, colors.white]),
                    ("BACKGROUND", (0, -1), (-1, -1), colors.lightblue),
                    ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ]
            )
        )
        elements.append(table)
    else:
        elements.append(
            Paragraph("Aucune donnee trouvee pour cette periode.", styles["Normal"])
        )


def _add_expenses_chart(elements, data):
    """Génère un camembert matplotlib des dépenses par catégorie et l'ajoute au PDF."""
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from reportlab.platypus import Image, Spacer

    items = data.get("items", [])
    if not items:
        return

    categories = [item["category"] for item in items]
    amounts = [item["amount"] for item in items]

    fig, ax = plt.subplots(1, 1, figsize=(6, 4))
    colors = plt.cm.Set3(range(len(categories)))
    wedges, texts, autotexts = ax.pie(
        amounts,
        labels=categories,
        autopct="%1.1f%%",
        colors=colors,
        textprops={"fontsize": 8},
    )
    ax.set_title("Repartition des depenses par categorie", fontsize=10)
    plt.tight_layout()

    img_buffer = BytesIO()
    fig.savefig(img_buffer, format="png", dpi=150, bbox_inches="tight")
    plt.close(fig)
    img_buffer.seek(0)

    elements.append(Spacer(1, 20))
    img = Image(img_buffer, width=400, height=270)
    elements.append(img)


def _add_income_vs_expenses_chart(elements, data):
    """Génère un graphique barres revenus/dépenses par période et l'ajoute au PDF."""
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from reportlab.platypus import Image, Spacer

    periods = data.get("periods", [])
    incomes = data.get("income", [])
    expenses = data.get("expenses", [])

    if not periods:
        return

    fig, ax = plt.subplots(1, 1, figsize=(7, 4))
    x = range(len(periods))
    width = 0.35

    ax.bar([i - width / 2 for i in x], incomes, width, label="Revenus", color="#10B981")
    ax.bar(
        [i + width / 2 for i in x], expenses, width, label="Depenses", color="#EF4444"
    )

    ax.set_xlabel("Periode")
    ax.set_ylabel("Montant (EUR)")
    ax.set_title("Revenus vs Depenses par mois", fontsize=10)
    ax.set_xticks(list(x))
    ax.set_xticklabels(periods, rotation=45, ha="right", fontsize=8)
    ax.legend()
    plt.tight_layout()

    img_buffer = BytesIO()
    fig.savefig(img_buffer, format="png", dpi=150, bbox_inches="tight")
    plt.close(fig)
    img_buffer.seek(0)

    elements.append(Spacer(1, 20))
    img = Image(img_buffer, width=430, height=250)
    elements.append(img)


# Export Excel : un classeur de plusieurs onglets reliés par des formules.
# « Opérations » contient le détail ; « Par catégorie », « Par mois » et
# « Synthèse » en sont calculés (SUMIFS) : corriger ou reclasser une opération
# dans Excel met tout le classeur à jour.
_XL_POLICE = "Arial"
_XL_BLEU = "1F4E79"
_XL_BLEU_CLAIR = "DDEBF7"
_XL_GRIS = "595959"
_XL_BANDE = "F2F2F2"
_XL_EUR = '#,##0.00 "€";[Red]-#,##0.00 "€";"-"'
_XL_PCT = '0.0%;[Red]-0.0%;"-"'
_XL_COULEURS = {"revenus": "10B981", "depenses": "EF4444", "solde": "1F4E79"}
_XL_MOIS = [
    "janvier", "février", "mars", "avril", "mai", "juin", "juillet",
    "août", "septembre", "octobre", "novembre", "décembre",
]  # fmt: skip
_XL_MAX_OPERATIONS = 20000

_XL_TITRES = {
    "expenses": "Dépenses par catégorie",
    "expenses_by_category": "Dépenses par catégorie",
    "expenses-by-category": "Dépenses par catégorie",
    "income-vs-expenses": "Revenus et dépenses",
    "cashflow": "Flux de trésorerie",
}


def _xl_type(type_transaction):
    """Type affiché dans l'onglet Opérations (critère des formules SUMIFS)."""
    if type_transaction in TransactionType.income_values():
        return "Revenu"
    if type_transaction in TransactionType.expense_values():
        return "Dépense"
    if type_transaction in TransactionType.transfer_values():
        return "Virement"
    return "Autre"


def _collect_export_data(user_id, start_date, end_date):
    """Données des exports Excel et JSON : dépenses par catégorie (ordre
    décroissant), mois de la période et opérations de la période."""
    from sqlalchemy.orm import aliased

    from core.models.account import Account

    Destination = aliased(Account)
    rows = (
        db.session.query(
            Transaction.id,
            Transaction.date_heure_transaction,
            Transaction.libelle,
            func.coalesce(Category.name, "Non catégorisée"),
            Account.name,
            Destination.name,
            Transaction.type_transaction,
            Transaction.montant,
            Transaction.notes,
        )
        .select_from(Transaction)
        .outerjoin(Category, Transaction.category_id == Category.id)
        .outerjoin(Account, Transaction.account_id == Account.id)
        .outerjoin(Destination, Transaction.destination_account_id == Destination.id)
        .filter(
            Transaction.user_id == user_id,
            func.date(Transaction.date_heure_transaction) >= start_date,
            func.date(Transaction.date_heure_transaction) <= end_date,
        )
        .order_by(Transaction.date_heure_transaction, Transaction.id)
        .limit(_XL_MAX_OPERATIONS)
        .all()
    )
    operations = [
        {
            "id": id_,
            "date": quand.date() if quand else None,
            "libelle": libelle or "",
            "categorie": categorie,
            "compte": compte or "",
            "destination": destination,
            "type": _xl_type(type_),
            "montant": abs(float(montant or 0)),
            "notes": notes,
        }
        for id_, quand, libelle, categorie, compte, destination, type_, montant, notes in rows
    ]

    mois = []
    annee, numero = start_date.year, start_date.month
    while (annee, numero) <= (end_date.year, end_date.month):
        mois.append((annee, numero))
        annee, numero = (annee + 1, 1) if numero == 12 else (annee, numero + 1)

    return {
        "categories": [
            i["category"]
            for i in _generate_expenses_report(user_id, start_date, end_date)["items"]
        ],
        "mois": mois,
        "operations": operations,
    }


def _xl_entete_feuille(ws, titre, sous_titre):
    from openpyxl.styles import Font

    ws.sheet_view.showGridLines = False
    ws["A1"] = titre
    ws["A1"].font = Font(name=_XL_POLICE, size=16, bold=True, color=_XL_BLEU)
    ws["A2"] = sous_titre
    ws["A2"].font = Font(name=_XL_POLICE, size=10, color=_XL_GRIS)
    ws.row_dimensions[1].height = 24


def _xl_tableau(ws, ligne, entetes, lignes, formats, total=None):
    """Tableau à partir de `ligne` : en-tête, lignes en bandes, puis TOTAL si
    `total` est fourni (fonction (lettre, première, dernière) -> formule ou
    None). Retourne (première ligne, dernière ligne, ligne du total)."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    trait = Side(style="thin", color="BFBFBF")
    for col, texte in enumerate(entetes, 1):
        cell = ws.cell(row=ligne, column=col, value=texte)
        cell.font = Font(name=_XL_POLICE, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color=_XL_BLEU)
        # Titre aligné comme sa colonne : texte à gauche, nombres à droite
        texte_ou_date = formats[col - 1] in ("@", "dd/mm/yyyy")
        cell.alignment = Alignment(
            horizontal="left" if texte_ou_date else "right", vertical="center"
        )
    ws.row_dimensions[ligne].height = 20

    premiere = ligne + 1
    for i, valeurs in enumerate(lignes):
        for col, valeur in enumerate(valeurs, 1):
            cell = ws.cell(row=premiere + i, column=col, value=valeur)
            cell.font = Font(name=_XL_POLICE)
            cell.number_format = formats[col - 1]
            cell.border = Border(bottom=trait)
            if i % 2:
                cell.fill = PatternFill("solid", start_color=_XL_BANDE)
    derniere = premiere + len(lignes) - 1

    r_total = None
    if total:
        r_total = derniere + 1
        for col in range(1, len(entetes) + 1):
            formule = (
                "TOTAL"
                if col == 1
                else total(get_column_letter(col), premiere, derniere)
            )
            cell = ws.cell(row=r_total, column=col, value=formule)
            cell.font = Font(name=_XL_POLICE, bold=True)
            cell.number_format = formats[col - 1]
            cell.border = Border(top=Side(style="thin"), bottom=Side(style="double"))
    return premiere, derniere, r_total


def _xl_largeurs(ws, largeurs):
    from openpyxl.utils import get_column_letter

    for col, largeur in enumerate(largeurs, 1):
        ws.column_dimensions[get_column_letter(col)].width = largeur


def _xl_impression(ws, titre, paysage=False, ligne_titres=None):
    """A4, largeur ajustée à la page, en-tête et pied de page."""
    ws.page_setup.orientation = "landscape" if paysage else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    if ligne_titres:
        ws.print_title_rows = f"{ligne_titres}:{ligne_titres}"
    ws.oddHeader.left.text = titre
    ws.oddFooter.left.text = "Gestion Budgétaire"
    ws.oddFooter.right.text = "Page &P / &N"


def _xl_feuille_operations(wb, operations, periode):
    """Détail des opérations, avec filtres. Retourne les plages
    (colonne -> 'Opérations'!$X$5:$X$n) utilisées par les formules."""
    from openpyxl.styles import Alignment, Font

    ws = wb.create_sheet("Opérations")
    _xl_entete_feuille(ws, "Opérations", periode)
    lignes = [
        [op["date"], op["libelle"], op["categorie"], op["compte"], op["type"],
         op["montant"]]  # fmt: skip
        for op in operations
    ]
    # Mise en forme directe (pas un style de tableau Excel : LibreOffice et
    # les aperçus ne l'affichent pas)
    _xl_tableau(
        ws,
        4,
        ["Date", "Libellé", "Catégorie", "Compte", "Type", "Montant"],
        lignes,
        ["dd/mm/yyyy", "@", "@", "@", "@", _XL_EUR],
    )
    derniere = 4 + max(len(operations), 1)
    if operations:
        ws.auto_filter.ref = f"A4:F{derniere}"
        for r in range(5, derniere + 1):
            ws.cell(row=r, column=1).alignment = Alignment(horizontal="center")
    else:
        ws["A5"] = "Aucune opération sur la période"
        ws["A5"].font = Font(name=_XL_POLICE, italic=True, color=_XL_GRIS)

    _xl_largeurs(ws, [13, 46, 36, 20, 11, 15])
    ws.freeze_panes = "A5"
    _xl_impression(ws, "Opérations", ligne_titres=4)
    return {
        lettre: f"'Opérations'!${lettre}$5:${lettre}${derniere}" for lettre in "ACEF"
    }


def _xl_feuille_categories(wb, categories, plages, periode, include_charts):
    from openpyxl.styles import Font

    ws = wb.create_sheet("Par catégorie")
    _xl_entete_feuille(ws, "Dépenses par catégorie", periode)
    if not categories:
        ws["A4"] = "Aucune dépense sur la période"
        ws["A4"].font = Font(name=_XL_POLICE, italic=True, color=_XL_GRIS)
        _xl_largeurs(ws, [34])
        return ws, None, None, None

    n_total = 5 + len(categories)
    somme = f'SUMIFS({plages["F"]},{plages["C"]},A{{r}},{plages["E"]},"Dépense")'
    lignes = [
        [
            nom,
            "=" + somme.format(r=r),
            f"=IF($B${n_total}=0,0,B{r}/$B${n_total})",
            f'=COUNTIFS({plages["C"]},A{r},{plages["E"]},"Dépense")',
        ]
        for r, nom in enumerate(categories, 5)
    ]
    premiere, derniere, r_total = _xl_tableau(
        ws,
        4,
        ["Catégorie", "Montant", "Part", "Opérations"],
        lignes,
        ["@", _XL_EUR, _XL_PCT, "0"],
        lambda c, a, b: f"=SUM({c}{a}:{c}{b})",
    )
    _xl_largeurs(ws, [34, 15, 9, 12])
    ws.freeze_panes = "A5"
    _xl_impression(ws, "Dépenses par catégorie", paysage=include_charts, ligne_titres=4)
    if include_charts:
        _add_expenses_excel_chart(ws, premiere, derniere, ws, "F4")
    return ws, premiere, derniere, r_total


def _xl_feuille_mois(wb, mois, plages, periode, include_charts):
    ws = wb.create_sheet("Par mois")
    _xl_entete_feuille(ws, "Revenus et dépenses par mois", periode)

    def somme(type_, annee, numero):
        suivant = (annee + 1, 1) if numero == 12 else (annee, numero + 1)
        return (
            f'SUMIFS({plages["F"]},{plages["E"]},"{type_}",'
            f'{plages["A"]},">="&DATE({annee},{numero},1),'
            f'{plages["A"]},"<"&DATE({suivant[0]},{suivant[1]},1))'
        )

    lignes = [
        [
            f"{_XL_MOIS[numero - 1]} {annee}",
            "=" + somme("Revenu", annee, numero),
            "=" + somme("Dépense", annee, numero),
            f"=B{r}-C{r}",
            f"=D{r}" if r == 5 else f"=E{r - 1}+D{r}",
        ]
        for r, (annee, numero) in enumerate(mois, 5)
    ]
    premiere, derniere, r_total = _xl_tableau(
        ws,
        4,
        ["Mois", "Revenus", "Dépenses", "Solde du mois", "Solde cumulé"],
        lignes,
        ["@", _XL_EUR, _XL_EUR, _XL_EUR, _XL_EUR],
        lambda c, a, b: f"=E{b}" if c == "E" else f"=SUM({c}{a}:{c}{b})",
    )
    _xl_largeurs(ws, [17, 15, 15, 16, 16])
    ws.freeze_panes = "A5"
    _xl_impression(ws, "Revenus et dépenses par mois", paysage=include_charts,
                   ligne_titres=4)  # fmt: skip
    if include_charts:
        _add_income_vs_expenses_excel_chart(ws, 4, derniere, ws, "G4")
    return ws, premiere, derniere, r_total


def _xl_feuille_synthese(ws, periode, mois, cat, plages, include_charts):
    """Chiffres clés et 5 plus fortes dépenses, calculés depuis les autres onglets."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    ws_mois, m_premiere, m_derniere, m_total = mois
    ws_cat, c_premiere, c_derniere, _ = cat
    _xl_entete_feuille(ws, "Synthèse", periode)

    indicateurs = [
        ("Revenus", f"='Par mois'!B{m_total}", _XL_EUR),
        ("Dépenses", f"='Par mois'!C{m_total}", _XL_EUR),
        ("Solde", "=B5-B6", _XL_EUR),
        ("Taux d'épargne", "=IF(B5=0,0,B7/B5)", _XL_PCT),
        ("Opérations", f'=COUNT({plages["A"]})', "0"),
    ]
    trait = Side(style="thin", color="BFBFBF")
    ws["A4"] = "Chiffres clés"
    ws["A4"].font = Font(name=_XL_POLICE, bold=True, size=12, color=_XL_BLEU)
    for i, (libelle, formule, fmt) in enumerate(indicateurs):
        r = 5 + i
        a, b = ws.cell(row=r, column=1, value=libelle), ws.cell(
            row=r, column=2, value=formule
        )
        a.font = Font(name=_XL_POLICE, color=_XL_GRIS)
        b.font = Font(name=_XL_POLICE, bold=True, size=12)
        b.number_format = fmt
        b.alignment = Alignment(horizontal="right")
        for cell in (a, b):
            cell.fill = PatternFill("solid", start_color=_XL_BLEU_CLAIR)
            cell.border = Border(bottom=trait)
        ws.row_dimensions[r].height = 20

    ws["A12"] = "Plus fortes dépenses"
    ws["A12"].font = Font(name=_XL_POLICE, bold=True, size=12, color=_XL_BLEU)
    if c_premiere:
        top = range(c_premiere, min(c_premiere + 5, c_derniere + 1))
        _xl_tableau(
            ws,
            13,
            ["Catégorie", "Montant", "Part"],
            [[f"='Par catégorie'!{c}{r}" for c in "ABC"] for r in top],
            ["@", _XL_EUR, _XL_PCT],
        )
    else:
        ws["A13"] = "Aucune dépense sur la période"
        ws["A13"].font = Font(name=_XL_POLICE, italic=True, color=_XL_GRIS)

    _xl_largeurs(ws, [34, 16, 9])
    _xl_impression(ws, "Synthèse", paysage=include_charts)
    if include_charts and m_derniere >= m_premiere:
        _add_income_vs_expenses_excel_chart(ws_mois, 4, m_derniere, ws, "E4")


def _build_excel_report(report_type, data, start_date, end_date, include_charts=False):
    """Classeur Excel du rapport : Synthèse, Par catégorie, Par mois, Opérations.

    data : résultat de _collect_export_data. L'onglet du type de rapport
    demandé est ouvert en premier. include_charts : graphiques Excel natifs,
    liés aux cellules (ils suivent les modifications faites dans le classeur).
    """
    from openpyxl import Workbook

    periode = (
        f"Période du {start_date.strftime('%d/%m/%Y')} "
        f"au {end_date.strftime('%d/%m/%Y')} · "
        f"généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}"
    )
    wb = Workbook()
    wb.properties.title = _XL_TITRES.get(report_type, "Rapport financier")
    wb.properties.creator = "Gestion Budgétaire"
    # Formules sans valeur en cache (openpyxl) : calcul à l'ouverture
    wb.calculation.fullCalcOnLoad = True

    # Pas de graphiques vides quand la période ne contient aucune opération
    include_charts = include_charts and bool(data["operations"])
    synthese = wb.active
    synthese.title = "Synthèse"
    plages = _xl_feuille_operations(wb, data["operations"], periode)
    cat = _xl_feuille_categories(
        wb, data["categories"], plages, periode, include_charts
    )
    mois = _xl_feuille_mois(wb, data["mois"], plages, periode, include_charts)
    _xl_feuille_synthese(synthese, periode, mois, cat, plages, include_charts)

    # Ordre de lecture : du résumé vers le détail
    wb.move_sheet("Par catégorie", offset=-1)
    wb.move_sheet("Par mois", offset=-1)
    if report_type in ("expenses", "expenses_by_category", "expenses-by-category"):
        wb.active = wb.sheetnames.index("Par catégorie")
    elif report_type in ("income-vs-expenses", "cashflow"):
        wb.active = wb.sheetnames.index("Par mois")
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws is wb.active

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _add_expenses_excel_chart(ws, first_row, last_row, cible, ancre):
    """Barres horizontales des dépenses par catégorie, la plus forte en haut.

    Plus lisible qu'un camembert dès qu'il y a plus de quelques catégories.
    """
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList

    chart = BarChart()
    chart.type = "bar"
    chart.title = "Dépenses par catégorie"
    chart.legend = None
    chart.add_data(
        Reference(ws, min_col=2, min_row=first_row, max_row=last_row),
        titles_from_data=False,
    )
    chart.set_categories(Reference(ws, min_col=1, min_row=first_row, max_row=last_row))
    chart.series[0].graphicalProperties.solidFill = _XL_BLEU
    # Montant au bout de chaque barre : pas besoin de le lire sur l'axe
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    # Sinon LibreOffice ajoute catégorie, série et symbole à chaque étiquette
    chart.dataLabels.showCatName = False
    chart.dataLabels.showSerName = False
    chart.dataLabels.showLegendKey = False
    chart.dataLabels.showPercent = False
    chart.dataLabels.position = "outEnd"
    chart.dataLabels.numFmt = "#,##0 €"
    # Le tableau est trié par montant décroissant : axe inversé pour le garder en haut
    chart.x_axis.scaling.orientation = "maxMin"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.y_axis.numFmt = "#,##0 €"
    chart.y_axis.majorGridlines = None
    chart.gapWidth = 40
    nb = last_row - first_row + 1
    chart.width, chart.height = 18, max(7, 0.6 * nb + 2)
    cible.add_chart(chart, ancre)


def _add_income_vs_expenses_excel_chart(ws, header_row, last_row, cible, ancre):
    """Barres revenus / dépenses par mois, avec le solde du mois en courbe."""
    from openpyxl.chart import BarChart, LineChart, Reference

    categories = Reference(ws, min_col=1, min_row=header_row + 1, max_row=last_row)

    bars = BarChart()
    bars.type = "col"
    bars.title = "Revenus et dépenses par mois"
    bars.add_data(
        Reference(ws, min_col=2, max_col=3, min_row=header_row, max_row=last_row),
        titles_from_data=True,
    )
    bars.set_categories(categories)
    bars.series[0].graphicalProperties.solidFill = _XL_COULEURS["revenus"]
    bars.series[1].graphicalProperties.solidFill = _XL_COULEURS["depenses"]
    bars.y_axis.numFmt = "#,##0 €"
    # Mois en bas du graphique, pas sur la ligne du zéro (soldes négatifs)
    bars.x_axis.tickLblPos = "low"
    bars.x_axis.delete = False
    bars.y_axis.delete = False
    bars.gapWidth = 60

    solde = LineChart()
    solde.add_data(
        Reference(ws, min_col=4, min_row=header_row, max_row=last_row),
        titles_from_data=True,
    )
    # Segments droits : une courbe lissée passe par des soldes qui n'existent pas
    solde.series[0].smooth = False
    solde.series[0].graphicalProperties.line.solidFill = _XL_COULEURS["solde"]
    solde.series[0].graphicalProperties.line.width = 28575  # 2,25 pt
    bars += solde

    bars.legend.position = "b"
    bars.width, bars.height = 22, 11
    cible.add_chart(bars, ancre)


def _build_csv_report(report_type, data):
    """Construit un fichier CSV pour un rapport."""
    output = StringIO()
    writer = csv.writer(output)

    if report_type in ("expenses", "expenses_by_category", "expenses-by-category"):
        writer.writerow(["Categorie", "Montant (EUR)", "Pourcentage (%)"])
        items = data.get("items", [])
        for item in items:
            writer.writerow([item["category"], item["amount"], item["percentage"]])
        writer.writerow(["TOTAL", data.get("total", 0), "100"])

    elif report_type in ("income-vs-expenses", "cashflow"):
        writer.writerow(["Periode", "Revenus (EUR)", "Depenses (EUR)", "Solde (EUR)"])
        periods = data.get("periods", [])
        incomes = data.get("income", [])
        expenses = data.get("expenses", [])
        for p, inc, exp in zip(periods, incomes, expenses):
            writer.writerow([p, inc, exp, inc - exp])
        writer.writerow(
            [
                "TOTAL",
                data.get("total_income", 0),
                data.get("total_expenses", 0),
                data.get("savings", 0),
            ]
        )

    return output.getvalue()


def _build_json_report(report_type, data, collected, start_date, end_date):
    """Export JSON complet : mêmes contenus que le classeur Excel (synthèse,
    catégories, mois, opérations). `data` (le rapport demandé) est conservé
    sous la clé « data » pour les lecteurs de l'ancien format."""
    operations = collected["operations"]

    def somme(ops):
        return round(sum(op["montant"] for op in ops), 2)

    revenus = [op for op in operations if op["type"] == "Revenu"]
    depenses = [op for op in operations if op["type"] == "Dépense"]
    total_revenus, total_depenses = somme(revenus), somme(depenses)
    solde = round(total_revenus - total_depenses, 2)

    by_category = []
    for nom in collected["categories"]:
        ops = [op for op in depenses if op["categorie"] == nom]
        montant = somme(ops)
        by_category.append(
            {
                "category": nom,
                "amount": montant,
                "percentage": (
                    round(montant / total_depenses * 100, 1) if total_depenses else 0
                ),
                "operations": len(ops),
            }
        )

    by_month, cumul = [], 0.0
    for annee, numero in collected["mois"]:

        def du_mois(ops):
            return [
                op
                for op in ops
                if op["date"] and (op["date"].year, op["date"].month) == (annee, numero)
            ]

        rev, dep = somme(du_mois(revenus)), somme(du_mois(depenses))
        cumul = round(cumul + rev - dep, 2)
        by_month.append(
            {
                "month": f"{annee}-{numero:02d}",
                "label": f"{_XL_MOIS[numero - 1]} {annee}",
                "income": rev,
                "expenses": dep,
                "balance": round(rev - dep, 2),
                "cumulative_balance": cumul,
            }
        )

    export = {
        "report": {
            "type": report_type,
            "title": _XL_TITRES.get(report_type, "Rapport financier"),
            "period": {"start": start_date.isoformat(), "end": end_date.isoformat()},
            "currency": "EUR",
            "generated_at": datetime.now().isoformat(timespec="seconds"),
        },
        "summary": {
            "total_income": total_revenus,
            "total_expenses": total_depenses,
            "balance": solde,
            "savings_rate": (
                round(solde / total_revenus * 100, 1) if total_revenus else 0
            ),
            "operations_count": len(operations),
        },
        "by_category": by_category,
        "by_month": by_month,
        "operations": [
            {
                "id": op["id"],
                "date": op["date"].isoformat() if op["date"] else None,
                "description": op["libelle"],
                "category": op["categorie"],
                "account": op["compte"],
                "destination_account": op["destination"],
                "type": {
                    "Revenu": "income",
                    "Dépense": "expense",
                    "Virement": "transfer",
                }.get(op["type"], "other"),
                "amount": round(op["montant"], 2),
                "notes": op["notes"],
            }
            for op in operations
        ],
        "data": data,
    }

    return json.dumps(export, ensure_ascii=False, indent=2, default=str)


# =============================
# ROUTES API
# =============================

# -- LIST -------------------------------------------------------------------


@api_v1_bp.route("/rapports", methods=["GET"])
@login_required
def api_reports_list():
    """Liste des rapports sauvegardés."""
    try:
        user_id = current_user.id
        from core.services.reports.report_services import ReportService

        reports = ReportService.get_user_reports(user_id)

        reports_data = [
            {
                "id": r.id,
                "name": r.name,
                "type": r.report_type,
                "parameters": r.parameters,
                "is_public": r.is_public,
                "created_at": r.created_at.isoformat(),
                "updated_at": r.updated_at.isoformat(),
            }
            for r in reports
        ]

        return api_success(reports_data)

    except Exception as e:
        logger.error(f"Erreur liste rapports: {e}", exc_info=True)
        return api_error("Une erreur interne est survenue", 500)


# -- GENERATE ---------------------------------------------------------------


@api_v1_bp.route("/rapports/generate", methods=["POST"])
@login_required
@auth_limiter.limit("10 per minute")
def api_reports_generate():
    """Génère un rapport avec les filtres spécifiés."""
    try:
        user_id = current_user.id
        data = request.get_json()
        if not data:
            return api_error("Données manquantes", 400)

        report_type = _normalize_report_type(data.get("report_type", "expenses"))
        date_range = data.get("date_range", "current_month")
        name = data.get("name")

        # Calculer les dates selon la période
        start_date, end_date = _resolve_date_range(date_range, data)

        # Générer selon le type
        if report_type in ("expenses", "expenses-by-category"):
            report_data = _generate_expenses_report(user_id, start_date, end_date)
        elif report_type == "income-vs-expenses":
            report_data = _generate_income_vs_expenses_report(
                user_id, start_date, end_date
            )
        elif report_type == "cashflow":
            report_data = _generate_cashflow_report(user_id, start_date, end_date)
        else:
            report_data = _generate_expenses_report(user_id, start_date, end_date)

        # Sauvegarder si un nom est fourni
        saved_report = None
        if name:
            from core.services.reports.report_services import ReportService

            saved_report = ReportService.save_report(
                user_id=user_id,
                name=name,
                report_type=report_type,
                parameters={
                    "date_range": date_range,
                    "start_date": start_date.isoformat(),
                    "end_date": end_date.isoformat(),
                },
            )

        return api_success(
            {
                "id": saved_report.id if saved_report else None,
                "report_type": report_type,
                "period": {
                    "start": start_date.isoformat(),
                    "end": end_date.isoformat(),
                },
                "data": report_data,
            }
        )

    except Exception as e:
        logger.error(f"Erreur génération rapport: {e}", exc_info=True)
        return api_error("Une erreur interne est survenue", 500)


# -- DATA (charts) ----------------------------------------------------------


@api_v1_bp.route("/rapports/data", methods=["GET"])
@login_required
def api_reports_data():
    """Données de rapport pour les graphiques frontend."""
    try:
        user_id = current_user.id
        date_range = request.args.get("date_range", "current_month")
        report_type = _normalize_report_type(
            request.args.get("report_type", "expenses")
        )

        start_date, end_date = _resolve_date_range(date_range, request.args)

        if report_type in ("expenses", "expenses-by-category"):
            data = _generate_expenses_report(user_id, start_date, end_date)
        elif report_type == "income-vs-expenses":
            data = _generate_income_vs_expenses_report(user_id, start_date, end_date)
        elif report_type == "cashflow":
            data = _generate_cashflow_report(user_id, start_date, end_date)
        else:
            data = _generate_expenses_report(user_id, start_date, end_date)

        return api_success(
            {
                "report_type": report_type,
                "period": {
                    "start": start_date.isoformat(),
                    "end": end_date.isoformat(),
                },
                "data": data,
            }
        )

    except Exception as e:
        logger.error(f"Erreur données rapport: {e}", exc_info=True)
        return api_error("Une erreur interne est survenue", 500)


# -- GET by ID --------------------------------------------------------------


@api_v1_bp.route("/rapports/<int:report_id>", methods=["GET"])
@login_required
def api_reports_get(report_id):
    """Récupère un rapport sauvegardé."""
    try:
        report = Report.query.filter_by(id=report_id, user_id=current_user.id).first()
        if not report:
            return api_error("Rapport non trouvé", 404)

        # Régénérer les données à partir des paramètres sauvegardés
        params = report.parameters or {}
        start_str = params.get("start_date")
        end_str = params.get("end_date")
        start_date = (
            datetime.fromisoformat(start_str).date()
            if start_str
            else datetime.now().date().replace(day=1)
        )
        end_date = (
            datetime.fromisoformat(end_str).date() if end_str else datetime.now().date()
        )

        normalized_type = _normalize_report_type(report.report_type)
        if normalized_type in ("expenses", "expenses-by-category"):
            data = _generate_expenses_report(current_user.id, start_date, end_date)
        elif normalized_type == "income-vs-expenses":
            data = _generate_income_vs_expenses_report(
                current_user.id, start_date, end_date
            )
        elif normalized_type == "cashflow":
            data = _generate_cashflow_report(current_user.id, start_date, end_date)
        else:
            data = _generate_expenses_report(current_user.id, start_date, end_date)

        return api_success(
            {
                "id": report.id,
                "name": report.name,
                "type": report.report_type,
                "parameters": report.parameters,
                "created_at": report.created_at.isoformat(),
                "data": data,
            }
        )

    except Exception as e:
        logger.error(f"Erreur récupération rapport {report_id}: {e}", exc_info=True)
        return api_error("Une erreur interne est survenue", 500)


# -- DELETE -----------------------------------------------------------------


@api_v1_bp.route("/rapports/<int:report_id>", methods=["DELETE"])
@login_required
def api_reports_delete(report_id):
    """Supprime un rapport sauvegardé."""
    try:
        report = Report.query.filter_by(id=report_id, user_id=current_user.id).first()
        if not report:
            return api_error("Rapport non trouvé", 404)

        db.session.delete(report)
        db.session.commit()

        return api_success(message="Rapport supprimé")

    except Exception as e:
        db.session.rollback()
        logger.error(f"Erreur suppression rapport {report_id}: {e}", exc_info=True)
        return api_error("Une erreur interne est survenue", 500)


# -- EXPORT MULTI-FORMAT ----------------------------------------------------


@api_v1_bp.route("/rapports/export", methods=["GET"])
@login_required
@auth_limiter.limit("10 per minute")
def api_reports_export():
    """Exporte un rapport au format spécifié (pdf, excel, csv, json)."""
    try:
        user_id = current_user.id
        export_format = request.args.get("format", "pdf").lower()
        include_charts = request.args.get("include_charts", "false").lower() == "true"
        date_range = request.args.get("date_range", "current_month")
        report_type = _normalize_report_type(
            request.args.get("report_type", "expenses")
        )

        start_date, end_date = _resolve_date_range(date_range, request.args)

        # Générer les données du rapport
        if report_type in ("expenses", "expenses-by-category"):
            report_data = _generate_expenses_report(user_id, start_date, end_date)
        elif report_type == "income-vs-expenses":
            report_data = _generate_income_vs_expenses_report(
                user_id, start_date, end_date
            )
        elif report_type == "cashflow":
            report_data = _generate_cashflow_report(user_id, start_date, end_date)
        else:
            report_data = _generate_expenses_report(user_id, start_date, end_date)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        if export_format == "pdf":
            pdf_buffer = _build_pdf_report(
                report_type,
                report_data,
                start_date,
                end_date,
                include_charts=include_charts,
            )
            response = make_response(pdf_buffer.getvalue())
            response.headers["Content-Type"] = "application/pdf"
            response.headers["Content-Disposition"] = (
                f"attachment; filename=rapport_{report_type}_{timestamp}.pdf"
            )

        elif export_format == "excel":
            excel_buffer = _build_excel_report(
                report_type,
                _collect_export_data(user_id, start_date, end_date),
                start_date,
                end_date,
                include_charts,
            )
            response = make_response(excel_buffer.getvalue())
            response.headers["Content-Type"] = (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            response.headers["Content-Disposition"] = (
                f"attachment; filename=rapport_{report_type}_{timestamp}.xlsx"
            )

        elif export_format == "csv":
            csv_content = _build_csv_report(report_type, report_data)
            response = make_response(csv_content)
            response.headers["Content-Type"] = "text/csv; charset=utf-8"
            response.headers["Content-Disposition"] = (
                f"attachment; filename=rapport_{report_type}_{timestamp}.csv"
            )

        elif export_format == "json":
            json_content = _build_json_report(
                report_type,
                report_data,
                _collect_export_data(user_id, start_date, end_date),
                start_date,
                end_date,
            )
            response = make_response(json_content)
            response.headers["Content-Type"] = "application/json; charset=utf-8"
            response.headers["Content-Disposition"] = (
                f"attachment; filename=rapport_{report_type}_{timestamp}.json"
            )

        else:
            return api_error(
                f"Format d'export non supporté: {export_format}. Formats valides: pdf, excel, csv, json",
                400,
            )

        logger.info(
            f"Export {export_format} rapport '{report_type}' pour user {user_id}"
        )
        return response

    except ImportError as e:
        logger.error(f"Module manquant pour l'export: {e}")
        return api_error(f"Export non disponible. Module manquant: {e}", 500)
    except Exception as e:
        logger.error(
            f"Erreur export {request.args.get('format', 'pdf')}: {e}", exc_info=True
        )
        return api_error("Une erreur interne est survenue", 500)


# -- EXPORT PDF (alias pour compatibilité) ----------------------------------


@api_v1_bp.route("/rapports/export/pdf", methods=["GET"])
@login_required
@auth_limiter.limit("10 per minute")
def api_reports_export_pdf():
    """Exporte un rapport au format PDF (alias de /export?format=pdf)."""
    try:
        user_id = current_user.id
        include_charts = request.args.get("include_charts", "false").lower() == "true"
        date_range = request.args.get("date_range", "current_month")
        report_type = _normalize_report_type(
            request.args.get("report_type", "expenses")
        )

        start_date, end_date = _resolve_date_range(date_range, request.args)

        # Générer les données du rapport
        if report_type in ("expenses", "expenses-by-category"):
            report_data = _generate_expenses_report(user_id, start_date, end_date)
        elif report_type == "income-vs-expenses":
            report_data = _generate_income_vs_expenses_report(
                user_id, start_date, end_date
            )
        elif report_type == "cashflow":
            report_data = _generate_cashflow_report(user_id, start_date, end_date)
        else:
            report_data = _generate_expenses_report(user_id, start_date, end_date)

        # Générer le PDF
        pdf_buffer = _build_pdf_report(
            report_type,
            report_data,
            start_date,
            end_date,
            include_charts=include_charts,
        )

        response = make_response(pdf_buffer.getvalue())
        response.headers["Content-Type"] = "application/pdf"
        response.headers["Content-Disposition"] = (
            f'attachment; filename=rapport_{report_type}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.pdf'
        )

        logger.info(f"Export PDF rapport '{report_type}' pour user {user_id}")
        return response

    except ImportError:
        logger.error("Module reportlab non disponible")
        return api_error("Export PDF non disponible. Module reportlab manquant.", 500)
    except Exception as e:
        logger.error(f"Erreur export PDF: {e}", exc_info=True)
        return api_error("Une erreur interne est survenue", 500)
