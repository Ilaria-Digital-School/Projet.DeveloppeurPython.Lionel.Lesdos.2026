"""Tests d'intégration pour l'API des rapports (/api/v1/rapports)."""

import json
from datetime import datetime
from decimal import Decimal

import pytest

pytestmark = pytest.mark.integration

# Période fixe : rend les tests indépendants de la date du jour
JUIN = {"date_range": "custom", "start_date": "2026-06-01", "end_date": "2026-06-30"}
JUIN_QS = "date_range=custom&start_date=2026-06-01&end_date=2026-06-30"


@pytest.fixture(autouse=True)
def _sans_rate_limit(monkeypatch):
    # generate/export sont limités à 10/min : sans ça la suite déclenche des 429
    from api.v1.auth import auth_limiter

    monkeypatch.setattr(auth_limiter, "enabled", False)


@pytest.fixture
def loisirs(db_session, user):
    from core.models.category import Category

    cat = Category(
        user_id=user.id, name="Loisirs", is_expense=True, category_type="expense"
    )
    db_session.add(cat)
    db_session.flush()
    return cat


@pytest.fixture
def operations(db_session, user, account, category, loisirs):
    """Juin 2026 : 3 dépenses catégorisées, 1 non catégorisée, 1 revenu ;
    mai 2026 : 1 revenu et 1 dépense ; juillet : 1 dépense hors période."""
    from core.models.transactions import Transaction

    def t(montant, type_, jour, category_id=None):
        return Transaction(
            user_id=user.id,
            account_id=account.id,
            category_id=category_id,
            montant=Decimal(montant),
            type_transaction=type_,
            libelle=f"{type_} {montant}",
            date_heure_transaction=datetime(*jour, 12, 0),
            statut="completed",
            devise="EUR",
        )

    db_session.add_all(
        [
            t("60.00", "achat", (2026, 6, 3), category.id),
            t("40.00", "ACHAT", (2026, 6, 10), category.id),
            t("25.00", "abonnement", (2026, 6, 15), loisirs.id),
            t("15.00", "achat", (2026, 6, 20)),  # sans catégorie
            t("2000.00", "salaire", (2026, 6, 28)),
            t("1800.00", "salaire", (2026, 5, 28)),
            t("300.00", "loyer", (2026, 5, 5)),
            t("999.00", "achat", (2026, 7, 1), category.id),  # hors juin
            t("500.00", "virement", (2026, 6, 12)),  # ni revenu ni dépense
        ]
    )
    db_session.flush()


def _data(client, report_type, qs=JUIN_QS):
    resp = client.get(f"/api/v1/rapports/data?report_type={report_type}&{qs}")
    assert resp.status_code == 200, resp.get_json()
    return resp.get_json()["data"]


class TestRapportsAuth:
    @pytest.mark.parametrize(
        "url", ["/api/v1/rapports", "/api/v1/rapports/data", "/api/v1/rapports/export"]
    )
    def test_requires_auth(self, client, url):
        assert client.get(url).status_code in (302, 401)


class TestRapportDepensesParCategorie:
    def test_totaux_par_categorie(self, auth_client, operations):
        data = _data(auth_client, "expenses_by_category")
        par_cat = {i["category"]: i["amount"] for i in data["items"]}
        assert par_cat["Alimentation"] == 100.0  # types 'achat' et 'ACHAT'
        assert par_cat["Loisirs"] == 25.0
        # Trié par montant décroissant
        assert data["categories"][0] == "Alimentation"

    def test_depenses_sans_categorie_comptees(self, auth_client, operations):
        data = _data(auth_client, "expenses_by_category")
        assert data["total"] == 140.0
        assert {i["category"]: i["amount"] for i in data["items"]}[
            "Non catégorisée"
        ] == 15.0

    def test_pourcentages(self, auth_client, operations):
        data = _data(auth_client, "expenses_by_category")
        assert sum(i["percentage"] for i in data["items"]) == pytest.approx(
            100, abs=0.2
        )

    def test_periode_respectee(self, auth_client, operations):
        data = _data(
            auth_client,
            "expenses_by_category",
            "date_range=custom&start_date=2026-07-01&end_date=2026-07-31",
        )
        assert data["total"] == 999.0

    def test_isolation_entre_utilisateurs(
        self, auth_client, operations, admin_user, db_session
    ):
        from core.models.account import Account
        from core.models.transactions import Transaction

        acc = Account(user_id=admin_user.id, name="Autre", balance=Decimal("0"))
        db_session.add(acc)
        db_session.flush()
        db_session.add(
            Transaction(
                user_id=admin_user.id,
                account_id=acc.id,
                montant=Decimal("5000.00"),
                type_transaction="achat",
                libelle="Pas à moi",
                date_heure_transaction=datetime(2026, 6, 5, 12),
                statut="completed",
                devise="EUR",
            )
        )
        db_session.flush()
        assert _data(auth_client, "expenses_by_category")["total"] == 140.0

    def test_sans_transactions(self, auth_client):
        data = _data(auth_client, "expenses_by_category")
        assert data["total"] == 0
        assert data["items"] == []


class TestRapportRevenusDepenses:
    def test_par_mois(self, auth_client, operations):
        data = _data(
            auth_client,
            "income_vs_expenses",
            "date_range=custom&start_date=2026-05-01&end_date=2026-06-30",
        )
        assert data["periods"] == ["2026-05", "2026-06"]
        assert data["income"] == [1800.0, 2000.0]
        assert data["expenses"] == [300.0, 140.0]
        assert data["total_income"] == 3800.0
        assert data["total_expenses"] == 440.0
        assert data["savings"] == 3360.0

    def test_virements_exclus(self, auth_client, operations):
        data = _data(auth_client, "income_vs_expenses")
        assert data["total_income"] == 2000.0
        assert data["total_expenses"] == 140.0

    def test_cashflow_cumule(self, auth_client, operations):
        data = _data(
            auth_client,
            "cashflow",
            "date_range=custom&start_date=2026-05-01&end_date=2026-06-30",
        )
        assert data["cumulative"] == [1500.0, 3360.0]


class TestRapportsSauvegardes:
    def test_generer_lister_consulter_supprimer(self, auth_client, operations):
        resp = auth_client.post(
            "/api/v1/rapports/generate",
            json={"report_type": "expenses_by_category", "name": "Juin", **JUIN},
        )
        assert resp.status_code == 200, resp.get_json()
        rid = resp.get_json()["id"]
        assert rid is not None

        liste = auth_client.get("/api/v1/rapports").get_json()["data"]
        assert [r["name"] for r in liste] == ["Juin"]

        detail = auth_client.get(f"/api/v1/rapports/{rid}")
        assert detail.status_code == 200
        assert detail.get_json()["data"]["total"] == 140.0

        assert auth_client.delete(f"/api/v1/rapports/{rid}").status_code == 200
        assert auth_client.get(f"/api/v1/rapports/{rid}").status_code == 404

    def test_generer_sans_nom_ne_sauvegarde_pas(self, auth_client, operations):
        resp = auth_client.post(
            "/api/v1/rapports/generate", json={"report_type": "cashflow", **JUIN}
        )
        assert resp.status_code == 200
        assert resp.get_json()["id"] is None
        assert auth_client.get("/api/v1/rapports").get_json()["data"] == []

    def test_rapport_d_un_autre_utilisateur_inaccessible(
        self, auth_client, admin_user, db_session
    ):
        from core.models.report import Report

        r = Report(
            user_id=admin_user.id, name="Privé", report_type="expenses", parameters={}
        )
        db_session.add(r)
        db_session.flush()
        assert auth_client.get(f"/api/v1/rapports/{r.id}").status_code == 404
        assert auth_client.delete(f"/api/v1/rapports/{r.id}").status_code == 404

    def test_rapport_public_d_un_autre_utilisateur_absent_de_la_liste(
        self, auth_client, admin_user, db_session
    ):
        from core.models.report import Report

        db_session.add(
            Report(
                user_id=admin_user.id,
                name="Public",
                report_type="expenses",
                parameters={},
                is_public=True,
            )
        )
        db_session.flush()
        assert auth_client.get("/api/v1/rapports").get_json()["data"] == []


class TestRapportsExport:
    @pytest.mark.parametrize(
        "fmt, content_type",
        [
            ("csv", "text/csv"),
            ("json", "application/json"),
            ("excel", "spreadsheetml"),
            ("pdf", "application/pdf"),
        ],
    )
    @pytest.mark.parametrize(
        "report_type", ["expenses_by_category", "income_vs_expenses"]
    )
    def test_formats(self, auth_client, operations, fmt, content_type, report_type):
        resp = auth_client.get(
            f"/api/v1/rapports/export?format={fmt}&report_type={report_type}&{JUIN_QS}"
        )
        assert resp.status_code == 200, resp.get_data(as_text=True)[:300]
        assert content_type in resp.headers["Content-Type"]
        assert "attachment" in resp.headers["Content-Disposition"]
        assert len(resp.data) > 0

    def test_export_json_contient_les_donnees(self, auth_client, operations):
        resp = auth_client.get(
            f"/api/v1/rapports/export?format=json&report_type=expenses_by_category&{JUIN_QS}"
        )
        contenu = json.loads(resp.get_data(as_text=True))
        assert "140" in json.dumps(contenu)

    @staticmethod
    def _excel(client, report_type, options=""):
        resp = client.get(
            f"/api/v1/rapports/export?format=excel&report_type={report_type}"
            f"&{JUIN_QS}{options}"
        )
        assert resp.status_code == 200
        return resp.data

    @staticmethod
    def _graphiques(contenu):
        """Types de graphiques de chaque graphique du .xlsx, triés."""
        import re
        import zipfile
        from io import BytesIO

        with zipfile.ZipFile(BytesIO(contenu)) as xlsx:
            return sorted(
                tuple(
                    sorted(
                        set(
                            re.findall(
                                r"<(?:c:)?(pieChart|barChart|lineChart)>",
                                xlsx.read(nom).decode(),
                            )
                        )
                    )
                )
                for nom in xlsx.namelist()
                if nom.startswith("xl/charts/chart")
            )

    @pytest.mark.parametrize(
        "report_type, onglet_ouvert",
        [
            ("expenses_by_category", "Par catégorie"),
            ("income_vs_expenses", "Par mois"),
            ("cashflow", "Par mois"),
        ],
    )
    def test_excel_classeur_de_plusieurs_onglets(
        self, auth_client, operations, report_type, onglet_ouvert
    ):
        from io import BytesIO

        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(self._excel(auth_client, report_type)))
        assert wb.sheetnames == ["Synthèse", "Par catégorie", "Par mois", "Opérations"]
        assert wb.active.title == onglet_ouvert

    def test_excel_avec_graphiques(self, auth_client, operations):
        contenu = self._excel(
            auth_client, "expenses_by_category", "&include_charts=true"
        )
        # Synthèse et Par mois : barres + solde ; Par catégorie : barres
        assert self._graphiques(contenu) == [
            ("barChart",),
            ("barChart", "lineChart"),
            ("barChart", "lineChart"),
        ]

    def test_excel_sans_option_sans_graphique(self, auth_client, operations):
        assert self._graphiques(self._excel(auth_client, "expenses_by_category")) == []

    def test_excel_avec_graphiques_sans_donnees(self, auth_client):
        # Période vide : pas de graphique vide, et pas d'erreur
        contenu = self._excel(
            auth_client, "expenses_by_category", "&include_charts=true"
        )
        assert self._graphiques(contenu) == []

    def test_excel_onglets_calcules_depuis_les_operations(
        self, auth_client, operations
    ):
        """Totaux en formules sur l'onglet Opérations : modifier une opération
        dans Excel met à jour catégories, mois et synthèse."""
        from io import BytesIO

        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(self._excel(auth_client, "expenses_by_category")))
        ops = wb["Opérations"]
        lignes = [r for r in ops.iter_rows(min_row=5, values_only=True) if r[0]]
        # Juin : 4 dépenses, 1 revenu, 1 virement (ni revenu ni dépense)
        assert sorted(r[4] for r in lignes) == [
            "Dépense", "Dépense", "Dépense", "Dépense", "Revenu", "Virement",
        ]  # fmt: skip

        cat = wb["Par catégorie"]
        assert cat["A5"].value == "Alimentation"  # plus forte dépense en tête
        assert cat["B5"].value.startswith("=SUMIFS('Opérations'!$F$5:$F$")
        assert '"Dépense"' in cat["B5"].value

        mois = wb["Par mois"]
        assert mois["A5"].value == "juin 2026"
        assert "DATE(2026,6,1)" in mois["B5"].value
        assert wb["Synthèse"]["B5"].value.startswith("='Par mois'!B")

    def test_export_json_complet(self, auth_client, operations):
        """Le JSON contient les mêmes contenus que le classeur Excel."""
        resp = auth_client.get(
            f"/api/v1/rapports/export?format=json&report_type=expenses_by_category"
            f"&{JUIN_QS}"
        )
        contenu = json.loads(resp.get_data(as_text=True))

        assert contenu["report"]["title"] == "Dépenses par catégorie"
        assert contenu["report"]["currency"] == "EUR"
        assert contenu["summary"] == {
            "total_income": 2000.0,
            "total_expenses": 140.0,
            "balance": 1860.0,
            "savings_rate": 93.0,
            "operations_count": 6,  # juin : 4 dépenses, 1 revenu, 1 virement
        }
        assert contenu["by_category"][0] == {
            "category": "Alimentation",
            "amount": 100.0,
            "percentage": 71.4,
            "operations": 2,
        }
        assert contenu["by_month"] == [
            {
                "month": "2026-06",
                "label": "juin 2026",
                "income": 2000.0,
                "expenses": 140.0,
                "balance": 1860.0,
                "cumulative_balance": 1860.0,
            }
        ]
        ops = contenu["operations"]
        assert len(ops) == 6
        assert set(ops[0]) == {
            "id", "date", "description", "category", "account",
            "destination_account", "type", "amount", "notes",
        }  # fmt: skip
        assert sorted(op["type"] for op in ops) == [
            "expense", "expense", "expense", "expense", "income", "transfer",
        ]  # fmt: skip
        # Ancien format conservé
        assert contenu["data"]["total"] == 140.0

    def test_format_invalide(self, auth_client):
        resp = auth_client.get("/api/v1/rapports/export?format=docx")
        assert resp.status_code == 400

    def test_alias_pdf(self, auth_client, operations):
        resp = auth_client.get(f"/api/v1/rapports/export/pdf?{JUIN_QS}")
        assert resp.status_code == 200
        assert resp.headers["Content-Type"] == "application/pdf"
